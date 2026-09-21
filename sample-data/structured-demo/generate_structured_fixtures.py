"""Generate the M3A.2 structured-data manual/automated test pack (SYNTHETIC data only).

Tenant: "beta" (see schemas/tenants/beta.yaml — it seeds ONE existing tenant custom field,
"Badge Colour"). Produces:

  01_employees.csv      employee master: core fields + an existing tenant custom field (Badge Colour)
                        + an UNKNOWN business field (T-Shirt Size -> custom-field proposal)
                        + an irrelevant legacy field (Legacy Payroll Code -> explicitly ignored)
  02_hr_details.xlsx    relational child sheets attached by Employee ID:
                        Vehicles (E100 has 2 vehicles; one EXACT duplicate row for E101;
                                  one CONFLICT: E100 TN70AB1234 as Motorcycle AND Scooter;
                                  E103 adds a 2nd vehicle to an employee that already exists in the target)
                        Education (multiple entries for E100), Emergency Contacts, Dependents, Addresses

Run:  python sample-data/structured-demo/generate_structured_fixtures.py
"""
from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).resolve().parent

EMPLOYEES = [
    ["Employee ID", "Full Name", "Work Email", "Department", "Hire Date", "Employment Type", "Designation",
     "Mobile Phone", "Badge Colour", "T-Shirt Size", "Legacy Payroll Code"],
    ["E100", "Priya Sharma", "priya.sharma@beta.example", "Engineering", "2021-02-15", "Full Time",
     "Senior Engineer", "+91 98765 43210", "Blue", "M", "P-1001"],
    ["E101", "Arjun Mehta", "arjun.mehta@beta.example", "Sales", "2020-08-03", "Full Time",
     "Account Manager", "+91 98765 43211", "Red", "L", "P-1002"],
    ["E102", "Meera Nair", "meera.nair@beta.example", "Finance", "2022-11-21", "Part Time",
     "Analyst", "+91 98765 43212", "Green", "XL", "P-1003"],
    ["E103", "Rahul Verma", "rahul.verma@beta.example", "Sales", "2019-06-17", "Full Time",
     "Analyst", "+91 98765 43213", "Blue", "M", "P-1004"],
]

SHEETS = {
    "Vehicles": [
        ["Employee ID", "Vehicle Type", "Registration Number"],
        ["E100", "Car", "TN70XY9876"],
        ["E100", "Motorcycle", "TN70AB1234"],
        ["E100", "Scooter", "TN70AB1234"],      # CONFLICT: same registration, different type
        ["E101", "Car", "KA01CD5678"],
        ["E101", "Car", "KA01CD5678"],          # EXACT DUPLICATE row -> collapses, provenance kept
        ["E103", "Car", "KA05ZZ0002"],          # target already has KA05ZZ0001 -> item ADDED
    ],
    "Education": [
        ["Employee ID", "Level", "Institution", "Score", "Year"],
        ["E100", "10th", "ABC School", "92.4", "2014"],
        ["E100", "B.E.", "XYZ College", "8.1", "2020"],
        ["E101", "12th", "PQR School", "88", "2016"],
        ["E102", "MBA", "LMN Institute", "3.7", "2021"],
    ],
    "Emergency Contacts": [
        ["Employee ID", "Contact Name", "Relationship", "Phone"],
        ["E100", "Anita Sharma", "Spouse", "+91 91234 56780"],
        ["E102", "Ravi Nair", "Father", "+91 91234 56781"],
    ],
    "Dependents": [
        ["Employee ID", "Dependent Name", "Relationship", "Date of Birth"],
        ["E100", "Aarav Sharma", "Child", "2018-05-10"],
        ["E101", "Kavya Mehta", "Child", "2020-09-01"],
    ],
    "Addresses": [
        ["Employee ID", "Address Type", "Address Line 1", "City", "State", "Postal Code", "Country"],
        ["E100", "Home", "12 MG Road", "Chennai", "Tamil Nadu", "600001", "India"],
        ["E101", "Home", "45 Brigade Road", "Bengaluru", "Karnataka", "560001", "India"],
    ],
}


def main() -> None:
    with (HERE / "01_employees.csv").open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(EMPLOYEES)
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in SHEETS.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
        # keep identifiers / codes as text so nothing is coerced
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.number_format = "@"
    wb.save(HERE / "02_hr_details.xlsx")
    print("wrote", HERE / "01_employees.csv", "and", HERE / "02_hr_details.xlsx")


if __name__ == "__main__":
    main()
