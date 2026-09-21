"""Final demo: run the FOUR clean fixtures through the REAL application workflow.

Drives the ACTUAL pipeline exactly as a user's upload does — normal workbook ingestion -> profile ->
deterministic mapping -> Groq where unresolved -> guardrail policy -> human review where necessary ->
preparation -> compare with target -> sync to the org-scoped mock target -> complete — NOT adapter-only
calls, inside ONE LangSmith trace per migration (`migration.<scenario>`). Four distinct named
organizations also exercise organization isolation (each migration's target writes are scoped to its
own organization).

The four fixtures under ``sample-data/final-demo/`` each contain EXACTLY ONE sheet (``Employees``), so
the driver uploads each workbook as-is through ordinary ingestion — it never selects or strips a sheet.

Scenarios (increasing difficulty), each a distinct organization:
    clean_no_review  04_clean_no_review_100x20.xlsx  CleanCo Pvt Ltd  (fully deterministic, no review)
    low_ambiguity    01_low_ambiguity_50x10.xlsx     Acme Pvt Ltd     (mostly rules + a little AI)
    medium_ambiguity 02_medium_ambiguity_20x10.xlsx  BCD Pvt Ltd      (mixed + one shared-email review)
    high_ambiguity   03_high_ambiguity_10x10.xlsx    XYZ Pvt Ltd      (all headers need real AI)

`clean_no_review` is fully target-aligned, so it reaches Compare + Sync with zero review and zero
semantic model calls — the agent does NOT call AI when deterministic logic is sufficient.

Usage (offline validation, no Groq — deterministic, safe, fast):
    LLM_PROVIDER=fake DATA_DIR=/tmp/demo4 backend/.venv/bin/python -m scripts.final_four_migrations
Live proof (real Groq + LangSmith, bounded, conservative — the recording run):
    LLM_PROVIDER=groq LLM_MAX_CONCURRENCY=1 LANGSMITH_TRACING=true \
      LANGSMITH_PROJECT=darwinbox-final-demo DATA_DIR=/tmp/demo4_live \
      backend/.venv/bin/python -m scripts.final_four_migrations --live

Runs are sequential with a spacing delay; the adapter honours Retry-After/backoff. The model is never
weakened and the fake adapter is never used for the live proof.
"""
from __future__ import annotations

import argparse
import asyncio
import json
import sys
import time
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "backend"))

from app.alias_audit import inventory_hash                      # noqa: E402
from app.mapping_rules import map_table_deterministic           # noqa: E402
from app.observability import job_metrics, migration_trace, stage_span  # noqa: E402
from app.profiling import ColumnProfile                          # noqa: E402
from app.runtime import AppContext                               # noqa: E402
from app.config import get_settings                              # noqa: E402
from app.schema_loader import get_target_schema                 # noqa: E402

FIX_DIR = REPO / "sample-data" / "final-demo"
# (scenario key, fixture filename, organization id, human org name)
SCENARIOS = [
    ("clean_no_review", "04_clean_no_review_100x20.xlsx", "cleanco-pvt-ltd", "CleanCo Pvt Ltd"),
    ("low_ambiguity", "01_low_ambiguity_50x10.xlsx", "acme-pvt-ltd", "Acme Pvt Ltd"),
    ("medium_ambiguity", "02_medium_ambiguity_20x10.xlsx", "bcd-pvt-ltd", "BCD Pvt Ltd"),
    ("high_ambiguity", "03_high_ambiguity_10x10.xlsx", "xyz-pvt-ltd", "XYZ Pvt Ltd"),
]

# --- Documented scripted human decisions (same rules as the M3F driver) ----------------------------
MAPPING_HINTS: dict[str, str] = {
    "Email": "work_email", "Joining Effective Date": "hire_date",
    "Worker Reference": "employee_id", "Legal Display Name": "full_name",
    "Corporate Mailbox": "work_email", "Service Commencement": "hire_date",
    "Identity Category": "gender", "Lifecycle State": "employment_status",
    "Engagement Class": "employment_type", "Org Function": "department",
    "Reports To Reference": "manager_employee_id", "Position Caption": "designation",
    "Start of Service": "hire_date", "Worker Category": "employment_type",
    "Current Standing": "employment_status", "Reporting Reference": "manager_employee_id",
    "Role Caption": "designation",
}
ENUM_DECISIONS: dict[str, dict[str, str]] = {
    "gender": {"Man": "male", "Woman": "female", "Does not disclose": "undisclosed",
               "male": "male", "female": "female", "non_binary": "non_binary"},
    "employment_status": {"Currently Employed": "active", "Away on Leave": "on_leave",
                          "Serving Notice": "notice_period", "Voluntarily Exited": "terminated",
                          "Dismissed": "terminated", "Working": "active", "Away": "on_leave",
                          "Resigning": "notice_period", "Separated": "terminated"},
    "employment_type": {"Permanent Staff": "full_time", "Reduced Hours": "part_time",
                        "External Contractor": "contractor", "Student Intern": "intern",
                        "Seasonal Worker": "temporary", "Core Employee": "full_time",
                        "Agency Contractor": "contractor", "Hourly Associate": "part_time",
                        "Fixed Term": "temporary"},
    "department": {"Product Engineering": "Engineering", "Revenue": "Sales",
                   "People & Culture": "People Operations", "Corporate Finance": "Finance"},
}
SPACING_SECONDS_DEFAULT = 60  # between live migrations, to respect the free Groq TPM budget

# Expected shape (data rows, columns) per fixture — asserted before any run (M3H §1).
EXPECTED_SHAPE = {
    "01_low_ambiguity_50x10.xlsx": (50, 10),
    "02_medium_ambiguity_20x10.xlsx": (20, 10),
    "03_high_ambiguity_10x10.xlsx": (10, 10),
    "04_clean_no_review_100x20.xlsx": (100, 20),
}


def _span_for(kind: str) -> str:
    return {"INGEST_FILE": "source.parse_stage", "MAP": "mapping", "RESUME_MAPPING": "mapping",
            "PREPARE": "preparation.normalize_validate", "RESUME_PREPARATION": "preparation.normalize_validate",
            "TARGET_RECONCILE": "target.reconcile", "RESUME_TARGET_REVIEW": "target.reconcile",
            "DELIVER_OP": "delivery.execute", "ROLLBACK_OP": "delivery.rollback"}.get(kind, kind.lower())


def _read_fixture(fname: str) -> tuple[bytes, int, int, list[str]]:
    """Read the fixture EXACTLY as uploaded (raw bytes, unmodified), plus (cols, rows, headers) for the
    report. Asserts the workbook contains exactly one 'Employees' sheet and the expected dimensions so a
    stale/wrong fixture can never slip into the demo (M3H §1). The driver ingests these raw bytes through
    the ordinary upload path — no sheet is selected or stripped."""
    import openpyxl
    path = FIX_DIR / fname
    data = path.read_bytes()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    assert wb.sheetnames == ["Employees"], f"{fname}: sheets={wb.sheetnames} (want exactly ['Employees'])"
    ws = wb["Employees"]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h) for h in rows[0]]
    cols, nrows = len(headers), len(rows) - 1
    wb.close()
    exp = EXPECTED_SHAPE.get(fname)
    assert exp is None or (nrows, cols) == exp, f"{fname}: shape {nrows}x{cols} (want {exp[0]}x{exp[1]})"
    return data, cols, nrows, headers


def _unseen_headers(headers: list[str]) -> list[str]:
    schema = get_target_schema()
    profs = [ColumnProfile(profile_id=f"c{i}", table_id="t", col_index=i, header=h, non_empty_count=10,
                           missing_count=0, distinct_count=10, observed_types={"text": 10},
                           format_indicators={}, samples=[]) for i, h in enumerate(headers)]
    res = {r.header: r.resolved for r in map_table_deterministic(profs, schema, table_name="Employees")}
    return [h for h in headers if not res.get(h, False)]


def _j(v):
    if isinstance(v, str):
        try:
            return json.loads(v)
        except Exception:
            return v
    return v


def _resolve_open_reviews(ctx, job_id: str, tally: dict) -> bool:
    """Make every pending human decision using the documented scripted rules."""
    db = ctx.db
    did = False
    for iss in db.get_issues(job_id, status="open"):
        cands = _j(iss.get("candidate_target_fields")) or []
        target = iss.get("proposed_target_field") or MAPPING_HINTS.get(iss["source_header"]) \
            or (cands[0] if cands else None)
        if iss.get("proposed_target_field"):
            resolution = {"action": "approve", "corrected_target": None, "reason": "consultant confirms best match", "note": None, "actor": "human"}
        else:
            resolution = {"action": "correct", "corrected_target": target, "reason": "consultant chooses target", "note": None, "actor": "human"}
        db.resolve_mapping_issue_and_enqueue(job_id, iss["id"], expected_version=iss["version"], resolution=resolution)
        tally["mapping_reviews"] = tally.get("mapping_reviews", 0) + 1
        did = True
    _REMAP_OK = {"blocked_provider", "mapping_complete", "preparation_complete", "reconciliation_complete",
                 "awaiting_record_review", "awaiting_target_review", "preparing_records", "error"}
    for p in db.get_custom_field_proposals(job_id, status="open"):
        hint = MAPPING_HINTS.get(p["source_header"])
        eff = ctx.effective_schema(p["tenant_id"])
        base = {"profile_id": p["profile_id"], "table_id": p["table_id"], "source_header": p["source_header"],
                "actor": "human", "method": "human", "note": None}
        audit_after = {"proposal_id": p["id"], "source_header": p["source_header"], "tenant_id": p["tenant_id"], "note": None}
        status = (db.get_job(job_id) or {}).get("status")
        rerun = "MAP" if status in _REMAP_OK else None
        if hint and eff.get(hint) is not None:
            resolution = {"action": "map_target", "target_path": hint, "note": None, "reason": "consultant maps to target", "actor": "human"}
            decision = dict(base, target_field=hint, status="corrected", destination_kind=eff.destination_kind(hint),
                            reason=f"consultant mapped '{p['source_header']}' -> {hint}")
            audit_after.update(target=hint, destination_kind=eff.destination_kind(hint), status="corrected")
            db.resolve_custom_field_proposal(job_id, p["id"], expected_version=p["version"], resolution=resolution,
                                             new_status="mapped_target", definition=None, decision=decision,
                                             audit_event="issue_resolved", audit_after=audit_after, rerun_kind=rerun)
            tally["proposals_mapped"] = tally.get("proposals_mapped", 0) + 1
        else:
            resolution = {"action": "ignore", "note": None, "reason": "not needed in target", "actor": "human"}
            decision = dict(base, target_field=None, status="ignored", destination_kind="IGNORED",
                            reason="source field explicitly ignored by consultant")
            audit_after.update(status="ignored", destination_kind="IGNORED")
            db.resolve_custom_field_proposal(job_id, p["id"], expected_version=p["version"], resolution=resolution,
                                             new_status="ignored", definition=None, decision=decision,
                                             audit_event="source_field_ignored", audit_after=audit_after, rerun_kind=rerun)
            tally["proposals_ignored"] = tally.get("proposals_ignored", 0) + 1
        did = True
    for iss in db.get_record_issues(job_id, status="open"):
        t = iss["issue_type"]
        aff = _j(iss.get("affected")) or {}
        opts = _j(iss.get("options")) or []
        scope = _j(iss.get("scope")) or {}
        res = {"action": None, "value": None, "convention": None, "value_map": None, "pivot": None,
               "scope": None, "candidate_id": None, "reason": "scripted demo decision", "note": None, "actor": "human"}
        if t == "shared_email":
            cands = aff.get("candidates") or []
            first = cands[0]
            res.update(action="correct", candidate_id=first["candidate_id"],
                       value=f"{(first.get('business_key') or 'emp').lower()}.unique@demo.example")
            tally["shared_email"] = tally.get("shared_email", 0) + 1
        elif t == "unknown_enum":
            field = iss.get("field")
            distinct = [d.get("value") for d in aff.get("distinct_values", [])] or scope.get("unmapped_values") or []
            decided = ENUM_DECISIONS.get(field, {})
            allowed = set(opts)
            vmap = {}
            for v in distinct:
                tv = decided.get(str(v)) or decided.get(v)
                if tv is None and allowed:
                    tv = next(iter(sorted(allowed)))
                if tv is not None:
                    vmap[str(v)] = tv
            res.update(action="map_values", value_map=vmap)
            tally["unknown_enum"] = tally.get("unknown_enum", 0) + 1
        elif t == "ambiguous_date":
            res.update(action="correct", value=(opts[0].get("iso") if opts else None))
            tally["ambiguous_date"] = tally.get("ambiguous_date", 0) + 1
        elif t in ("value_conflict",):
            res.update(action="select", value=str(opts[0]) if opts else None)
            tally["value_conflict"] = tally.get("value_conflict", 0) + 1
        elif t == "orphan_child_row":
            res.update(action="exclude")
            tally["orphan"] = tally.get("orphan", 0) + 1
        else:
            res.update(action="exclude")
            tally["other_record"] = tally.get("other_record", 0) + 1
        db.resolve_record_issue_and_enqueue(job_id, iss["id"], expected_version=iss["version"], resolution=res)
        did = True
    for iss in db.get_target_review_issues(job_id, status="open"):
        res = {"action": "use_incoming", "reason": "scripted demo decision", "note": None, "actor": "human"}
        db.resolve_target_issue_and_enqueue(job_id, iss["id"], expected_version=iss["version"], resolution=res)
        tally["target_reviews"] = tally.get("target_reviews", 0) + 1
        did = True
    return did


async def _run_scenario(ctx, scenario: str, fname: str, tenant: str, org_name: str, project: str, live: bool) -> dict:
    db = ctx.db
    data, cols, rows, headers = _read_fixture(fname)
    unseen = _unseen_headers(headers)
    job_id = db.create_job(schema_version=ctx.schema.version, provider=ctx.provider,
                           model_id=ctx.model_id, adapter_kind=ctx.adapter_kind, tenant_id=tenant)
    info = ctx.blobstore.put(data, suffix=".xlsx")
    file_id = f"file_{info.key.split('.')[0][:12]}"
    db.add_source_file(job_id, file_id=file_id, original_filename=fname, stored_name=info.key,
                       content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       size_bytes=info.size_bytes, blob_key=info.key, sha256=info.sha256,
                       storage_status="stored", parse_status="queued")
    db.enqueue_work(job_id=job_id, kind="INGEST_FILE", source_file_id=file_id,
                    idempotency_key=f"{job_id}:ingest:{file_id}", max_attempts=ctx.settings.work_max_attempts)
    db.set_job_stage(job_id, status="queued", stage="queued")

    pool = ctx.workers
    worker_id = f"demo-{scenario}"
    tally: dict = {}
    trace_id = None
    t0 = time.time()
    with migration_trace(scenario, job_id, project_name=project, enabled=live) as rt:
        trace_id = getattr(rt, "id", None)
        guard = 0
        while True:
            guard += 1
            if guard > 800:
                break
            item = db.claim_next_work(worker_id, lease_seconds=ctx.settings.work_lease_seconds)
            if item is not None:
                with stage_span(_span_for(item["kind"]), enabled=live,
                                inputs={"job_id": job_id, "kind": item["kind"]}):
                    await pool._process(item, worker_id)
                continue
            status = (db.get_job(job_id) or {}).get("status")
            if status in ("awaiting_review", "blocked_provider", "awaiting_record_review",
                          "awaiting_target_review", "stale_target_review_required"):
                with stage_span("review.required", enabled=live, inputs={"status": status,
                                "open_mapping": len(db.get_issues(job_id, status="open")),
                                "open_record": len(db.get_record_issues(job_id, status="open")),
                                "open_target": len(db.get_target_review_issues(job_id, status="open")),
                                "open_proposals": len(db.get_custom_field_proposals(job_id, status="open"))}):
                    pass
                if _resolve_open_reviews(ctx, job_id, tally):
                    continue
            break
    wall = round(time.time() - t0, 1)

    m = job_metrics(db, job_id)
    final = (db.get_job(job_id) or {}).get("status")
    return {
        "scenario": scenario, "job_id": job_id, "fixture": fname, "tenant": tenant, "org_name": org_name,
        "dimensions": f"{rows}x{cols}", "unseen_headers": unseen,
        "final_status": final, "wall_seconds": wall, "review_tally": tally,
        "mapping": m.get("mapping", {}), "intelligence": m.get("intelligence", {}),
        "records": m.get("preparation", {}), "reconciliation": m.get("reconciliation", {}),
        "delivery": m.get("delivery", {}), "model": {k: v for k, v in (m.get("model") or {}).items() if k != "calls_detail"},
        "timings": m.get("timings", {}), "human_decisions": m.get("human_decisions", 0),
        "langsmith_trace": {"name": f"migration.{scenario}", "id": str(trace_id) if trace_id else None},
    }


async def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--live", action="store_true", help="real Groq + LangSmith (default: offline validation)")
    ap.add_argument("--spacing", type=int, default=None, help="seconds between migrations")
    args = ap.parse_args()

    settings = get_settings()
    live = args.live
    project = settings.langsmith_project
    spacing = args.spacing if args.spacing is not None else (SPACING_SECONDS_DEFAULT if live else 0)

    hash_before = inventory_hash(get_target_schema())
    ctx = await AppContext.create(settings, start_workers=False)
    print(f"adapter={ctx.adapter_kind} provider={ctx.provider} tracing={settings.tracing_enabled} "
          f"project={project} live={live}")
    results = []
    try:
        for i, (scenario, fname, tenant, org_name) in enumerate(SCENARIOS):
            print(f"\n=== {scenario} ({fname}) org={org_name} ===")
            r = await _run_scenario(ctx, scenario, fname, tenant, org_name, project, live)
            print(f"  final={r['final_status']} wall={r['wall_seconds']}s dims={r['dimensions']} "
                  f"mappings(rule/model/human)={r['mapping'].get('rule')}/{r['mapping'].get('model')}/{r['mapping'].get('human')} "
                  f"model_calls={r['model'].get('calls')} tokens={r['model'].get('total_tokens')} "
                  f"reviews={r['review_tally']}")
            results.append(r)
            if spacing and i < len(SCENARIOS) - 1:
                print(f"  spacing {spacing}s (Groq TPM budget)…")
                time.sleep(spacing)
    finally:
        await ctx.aclose()

    hash_after = inventory_hash(get_target_schema())
    _write_report(results, hash_before, hash_after, live, project, ctx.adapter_kind)
    ok = all(r["final_status"] == "migration_complete" for r in results)
    clean = next((r for r in results if r["scenario"] == "clean_no_review"), None)
    clean_zero_calls = clean is not None and (clean["model"].get("calls") or 0) == 0
    print(f"\nalias hash unchanged: {hash_before == hash_after}")
    print(f"all migration_complete: {ok}")
    print(f"clean_no_review model calls: {clean['model'].get('calls') if clean else '?'} (want 0)")
    return 0 if ok and hash_before == hash_after else 1


def _write_report(results, hash_before, hash_after, live, project, adapter_kind) -> None:
    out = REPO / "reports" / "final_four_migration_demo.md"
    out.parent.mkdir(exist_ok=True)
    clean = next((r for r in results if r["scenario"] == "clean_no_review"), None)
    lines = ["# Final four-migration demo (live Groq + LangSmith)", "",
             f"- mode: {'LIVE (real Groq + LangSmith)' if live else 'OFFLINE validation (fake adapter)'}",
             f"- adapter: `{adapter_kind}`", f"- LangSmith project: `{project}`",
             "- each workbook contains a single `Employees` sheet, uploaded as-is through ordinary ingestion (no sheet is selected or stripped)",
             f"- alias inventory hash before: `{hash_before}`",
             f"- alias inventory hash after: `{hash_after}`",
             f"- alias hash unchanged: **{hash_before == hash_after}**",
             f"- all reached migration_complete: **{all(r['final_status'] == 'migration_complete' for r in results)}**",
             f"- clean_no_review semantic model calls: **{(clean['model'].get('calls') if clean else '?')}** (want 0)",
             f"- four distinct organizations (isolation): {', '.join(r['org_name'] for r in results)}", ""]
    for r in results:
        mp, md, tim = r["mapping"], r["model"], r["timings"]
        lines += [f"## {r['scenario']} — `{r['fixture']}`", "",
                  f"- job id: `{r['job_id']}` · organization: **{r['org_name']}** (`{r['tenant']}`) · dimensions: {r['dimensions']}",
                  f"- **final status: {r['final_status']}** · wall {r['wall_seconds']}s",
                  f"- intended unseen headers ({len(r['unseen_headers'])}): {', '.join(r['unseen_headers']) or '(none — fully canonical)'}",
                  f"- field mapping — rule: {mp.get('rule')} · AI+guardrails: {mp.get('model')} · confirmed by user: {mp.get('human')} · not mapped: {mp.get('unmapped')} · not migrated: {mp.get('ignored')}",
                  f"- transforms — deterministic: {r['intelligence'].get('transforms_deterministic')} · model: {r['intelligence'].get('transforms_model')} · human: {r['intelligence'].get('transforms_human')}",
                  f"- reviews resolved (by type): {json.dumps(r['review_tally']) if r['review_tally'] else '{} (none)'}",
                  f"- Groq calls: {md.get('calls')} · attempts: {md.get('total_attempts')} · tokens (in/out/total): {md.get('prompt_tokens')}/{md.get('completion_tokens')}/{md.get('total_tokens')} · avg latency: {md.get('latency_ms_avg')} ms",
                  f"- stage timings (ms): {json.dumps(tim.get('stages', {}))} · total compute: {tim.get('total_compute_ms')} ms",
                  f"- records — eligible: {r['records'].get('eligible')} · blocked: {r['records'].get('blocked')} · excluded: {r['records'].get('excluded')}",
                  f"- compare — new: {r['reconciliation'].get('ready_create')} · update: {r['reconciliation'].get('ready_update')} · no-change: {r['reconciliation'].get('no_change')}",
                  f"- sync — {json.dumps(r['delivery'].get('by_status', {}))} · attempts: {r['delivery'].get('total_attempts')}",
                  f"- LangSmith trace: `{r['langsmith_trace']['name']}` id=`{r['langsmith_trace']['id']}`", ""]
    out.write_text("\n".join(lines))
    print(f"\nreport: {out}")


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
