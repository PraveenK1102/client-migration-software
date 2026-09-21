"""The four final demo fixtures under sample-data/final-demo/.

Guards the demo inputs so the recorded demo (and the demo driver) stay honest:
  * each fixture is a single ``Employees`` sheet at the intended shape (50x10 / 20x10 / 10x10 / 100x20)
    — so the driver ingests it through the ordinary upload path with no sheet selection;
  * every INTENDED unseen header is genuinely still unresolved by ``map_table_deterministic`` (no fixture
    header leaked into the deterministic alias surface — real semantic mapping is exercised);
  * the canonical headers still resolve deterministically (the clean file resolves ALL 20);
  * the medium fixture carries exactly one controlled shared-email collision;
  * the low/high/clean fixtures have no accidental shared-email collision;
  * the frozen alias-inventory hash is unchanged (no demo-specific alias was added).
"""
from __future__ import annotations

from collections import Counter
from pathlib import Path

import openpyxl
import pytest

from app.alias_audit import inventory_hash, resolves_deterministically
from app.mapping_rules import map_table_deterministic
from app.profiling import ColumnProfile
from app.schema_loader import get_target_schema
from tests.test_m3e_anti_overfitting import FROZEN_SURFACE_HASH

DEMO_DIR = Path(__file__).resolve().parents[2] / "sample-data" / "final-demo"

LOW = "01_low_ambiguity_50x10.xlsx"
MEDIUM = "02_medium_ambiguity_20x10.xlsx"
HIGH = "03_high_ambiguity_10x10.xlsx"
CLEAN = "04_clean_no_review_100x20.xlsx"

# Expected (data rows, columns) per fixture on the Employees sheet.
SHAPE = {LOW: (50, 10), MEDIUM: (20, 10), HIGH: (10, 10), CLEAN: (100, 20)}

# Headers that MUST require semantic (model) mapping — never a deterministic alias.
UNSEEN = {
    LOW: {"Email", "Joining Effective Date"},
    MEDIUM: {"Start of Service", "Worker Category", "Current Standing", "Reporting Reference",
             "Role Caption"},
    HIGH: {"Worker Reference", "Legal Display Name", "Corporate Mailbox", "Service Commencement",
           "Identity Category", "Lifecycle State", "Engagement Class", "Org Function",
           "Reports To Reference", "Position Caption"},
    CLEAN: set(),
}
# Headers that should still resolve by the frozen deterministic rule.
CANONICAL = {
    LOW: {"employee_id", "full_name", "gender", "employment_status", "employment_type",
          "department", "manager_employee_id", "designation"},
    MEDIUM: {"employee_id", "full_name", "work_email", "gender", "department"},
    HIGH: set(),
    # The clean happy-path file is fully target-aligned: EVERY header resolves deterministically.
    CLEAN: {"employee_id", "full_name", "preferred_name", "date_of_birth", "gender", "nationality",
            "hire_date", "contract_start_date", "employment_type", "employment_status", "designation",
            "grade", "probation_end_date", "notice_period_days", "department", "business_unit",
            "cost_center", "work_location", "manager_employee_id", "work_email"},
}
ALL = list(SHAPE)


def _headers_and_rows(fname):
    wb = openpyxl.load_workbook(DEMO_DIR / fname, read_only=True, data_only=True)
    sheetnames = list(wb.sheetnames)
    ws = wb["Employees"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    return sheetnames, list(rows[0]), rows[1:]


@pytest.mark.parametrize("fname", ALL)
def test_fixture_is_a_single_employees_sheet_at_the_intended_shape(fname):
    """Each final workbook is ONE sheet named 'Employees' at the intended dimensions, so the demo
    driver uploads it as-is (no sheet is selected or stripped)."""
    sheetnames, headers, data = _headers_and_rows(fname)
    assert sheetnames == ["Employees"], f"{fname}: sheets={sheetnames} (want exactly ['Employees'])"
    exp_rows, exp_cols = SHAPE[fname]
    assert len(headers) == exp_cols, f"{fname}: {len(headers)} columns (want {exp_cols})"
    assert len(data) == exp_rows, f"{fname}: {len(data)} data rows (want {exp_rows})"


@pytest.mark.parametrize("fname", ALL)
def test_intended_unseen_headers_are_unresolved_deterministically(fname):
    _, headers, _ = _headers_and_rows(fname)
    schema = get_target_schema()
    results = {r.header: r for r in map_table_deterministic(_profiles(headers), schema, table_name="Employees")}
    for h in UNSEEN[fname]:
        assert h in results, f"{fname}: header {h!r} not in fixture"
        assert results[h].resolved is False, f"{fname}: {h!r} resolved deterministically to {results[h].target}"
        # the standalone probe agrees (no alias was added for the fixture)
        assert resolves_deterministically(schema, h) is None, f"{fname}: {h!r} is a deterministic alias"


@pytest.mark.parametrize("fname", ALL)
def test_canonical_headers_still_resolve(fname):
    _, headers, _ = _headers_and_rows(fname)
    schema = get_target_schema()
    results = {r.header: r for r in map_table_deterministic(_profiles(headers), schema, table_name="Employees")}
    for h in CANONICAL[fname]:
        assert results[h].resolved is True, f"{fname}: canonical {h!r} did not resolve"


def test_clean_fixture_resolves_every_header_deterministically():
    """The happy-path file must map fully by rules — no unresolved header, so it can reach Compare with
    target and Sync WITHOUT any semantic model call for field mapping."""
    _, headers, _ = _headers_and_rows(CLEAN)
    schema = get_target_schema()
    results = {r.header: r for r in map_table_deterministic(_profiles(headers), schema, table_name="Employees")}
    unresolved = [h for h in headers if not results[h].resolved]
    assert unresolved == [], f"clean fixture has unresolved headers: {unresolved}"


def test_medium_fixture_has_the_controlled_shared_email_collision():
    """Exactly two employees share shared.medium@demo.example (one clear shared-email review)."""
    _, headers, data = _headers_and_rows(MEDIUM)
    hi = {h: i for i, h in enumerate(headers)}
    by_email = Counter(row[hi["work_email"]] for row in data)
    shared = {e: c for e, c in by_email.items() if c > 1}
    assert shared == {"shared.medium@demo.example": 2}, shared


def test_low_high_clean_fixtures_have_no_duplicate_email_source():
    """Fixtures 1, 3 and 4 must NOT trigger an accidental shared-email review."""
    for fname, email_hdr in [(LOW, "Email"), (HIGH, "Corporate Mailbox"), (CLEAN, "work_email")]:
        _, headers, data = _headers_and_rows(fname)
        hi = {h: i for i, h in enumerate(headers)}
        emails = [row[hi[email_hdr]] for row in data]
        assert len(emails) == len(set(emails)), f"{fname}: duplicate emails present"


def test_demo_fixtures_do_not_change_the_alias_inventory_hash():
    """The deterministic alias surface is frozen; no fixture-specific alias was added."""
    assert inventory_hash(get_target_schema()) == FROZEN_SURFACE_HASH


def _profiles(headers):
    return [ColumnProfile(profile_id=f"c{i}", table_id="t", col_index=i, header=h,
                          non_empty_count=10, missing_count=0, distinct_count=10,
                          observed_types={"text": 10}, format_indicators={}, samples=[])
            for i, h in enumerate(headers)]
