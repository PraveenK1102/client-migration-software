"""M3D P0: NO silent row truncation.

Invariant: the system ingests the COMPLETE accepted table, or REJECTS the file. It never
migrates a silently-truncated subset of the employee population. Rejection is detected while
reading (at limit+1), no partial rows are persisted, and mapping/preparation/delivery never run
from a rejected source. Failure is clearly visible (file error + job error + audit).
"""
from __future__ import annotations

import io
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.ingest import parse_file
from app.source_records import RowLimitExceededError


def _csv_bytes(n_rows: int) -> bytes:
    lines = ["EmployeeID,Full Name,Email"]
    for i in range(1, n_rows + 1):
        lines.append(f"E{i:05d},Person {i},person{i}@example.com")
    return ("\n".join(lines) + "\n").encode()


def _xlsx_bytes(n_rows: int) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"
    ws.append(["EmployeeID", "Full Name", "Email"])
    for i in range(1, n_rows + 1):
        ws.append([f"E{i:05d}", f"Person {i}", f"person{i}@example.com"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _parse(name: str, data: bytes, max_rows: int):
    return parse_file(filename=name, data=data, stored_name="s_" + name,
                      max_bytes=50_000_000, max_rows=max_rows)


# ---------------------------------------------------------------- direct parser (CSV) ----------
def test_csv_exactly_at_limit_succeeds_with_all_rows():
    pf = _parse("at_limit.csv", _csv_bytes(100), max_rows=100)
    assert pf.tables[0].n_rows == 100
    assert len(pf.records) == 100
    # first + last rows are genuinely present (nothing dropped at either end)
    assert pf.records[0].cells[0].value == "E00001"
    assert pf.records[-1].cells[0].value == "E00100"


def test_csv_limit_plus_one_is_rejected_no_partial_rows():
    with pytest.raises(RowLimitExceededError) as ei:
        _parse("over.csv", _csv_bytes(101), max_rows=100)
    msg = str(ei.value)
    assert "REJECTED" in msg and "100" in msg
    # The exception carries no ParsedFile -> nothing partial can be persisted by the caller.


def test_csv_trailing_blank_lines_do_not_trip_the_limit():
    # A file of exactly `max_rows` real employees plus trailing blank lines must SUCCEED
    # (blank rows are noise, not a truncated employee, and are never counted or persisted).
    data = _csv_bytes(50) + b"\n\n\n"
    pf = _parse("trailing.csv", data, max_rows=50)
    assert pf.tables[0].n_rows == 50 and len(pf.records) == 50


# ---------------------------------------------------------------- direct parser (XLSX) ----------
def test_xlsx_exactly_at_limit_succeeds():
    pf = _parse("at_limit.xlsx", _xlsx_bytes(100), max_rows=100)
    assert pf.tables[0].n_rows == 100
    assert len(pf.records) == 100
    assert pf.records[-1].cells[0].value == "E00100"


def test_xlsx_limit_plus_one_is_rejected_no_partial_rows():
    with pytest.raises(RowLimitExceededError) as ei:
        _parse("over.xlsx", _xlsx_bytes(101), max_rows=100)
    assert "REJECTED" in str(ei.value)


# ---------------------------------------------------------------- stress: 10k rows -------------
@pytest.mark.parametrize("n", [10_000])
def test_generated_10k_rows_ingests_completely_with_raised_limit(n):
    """Stress (NOT a default-limit change): with an explicitly raised limit, a 10k-row file ingests
    COMPLETELY — proving completeness is enforced, not just the rejection path."""
    pf = _parse("big.csv", _csv_bytes(n), max_rows=n)
    assert pf.tables[0].n_rows == n and len(pf.records) == n
    assert pf.records[-1].cells[0].value == f"E{n:05d}"


# ---------------------------------------------------------------- end-to-end job behavior -------
@pytest.fixture
def small_limit_client(tmp_path, monkeypatch):
    """A real app whose hard row limit is small, so oversized-source behavior is cheap to exercise."""
    monkeypatch.setenv("LLM_PROVIDER", "fake")
    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("AUTO_CONTINUE", "false")
    monkeypatch.setenv("MAX_ROWS_PER_TABLE", "20")
    monkeypatch.delenv("GROQ_API_KEY", raising=False)
    from app import config
    config.get_settings.cache_clear()
    config.get_settings().ensure_dirs()
    from app.main import create_app
    app = create_app()
    with TestClient(app) as c:
        yield c
    config.get_settings.cache_clear()


def _poll(client, job_id, want, tries=120, delay=0.1):
    for _ in range(tries):
        j = client.get(f"/api/jobs/{job_id}").json()
        if j["status"] in want:
            return j
        time.sleep(delay)
    return client.get(f"/api/jobs/{job_id}").json()


def test_multi_file_job_with_one_oversized_source_fails_and_does_not_proceed(small_limit_client):
    client = small_limit_client
    files = [
        ("files", ("good.csv", _csv_bytes(20), "text/csv")),          # exactly at the limit -> OK
        ("files", ("oversized.csv", _csv_bytes(21), "text/csv")),      # limit+1 -> rejected
    ]
    r = client.post("/api/jobs", files=files)
    assert r.status_code == 200
    job_id = r.json()["id"]

    job = _poll(client, job_id, {"error"})
    # (5) the job MUST NOT proceed as though complete: it is a clear error, never mapping/etc.
    assert job["status"] == "error", job
    assert job.get("stage") == "ingest_failed"
    assert job["status"] not in ("mapping_complete", "awaiting_review", "preparation_complete")

    # (6) the failure is clearly reported: job error text, the specific file's parse_error, audit.
    assert "oversized.csv" in (job.get("error") or "")
    assert job["counts"]["files_failed"] >= 1
    files_state = client.get(f"/api/jobs/{job_id}/files").json()
    oversized = [f for f in files_state if f["original_filename"] == "oversized.csv"]
    assert oversized and oversized[0]["parse_status"] == "failed"
    assert "REJECTED" in (oversized[0].get("parse_error") or "")

    audit = client.get(f"/api/jobs/{job_id}/audit").json()
    assert any(a["event_type"] == "ingest_failed" for a in audit)

    # No mapping ever happened (no decisions), i.e. nothing proceeded from the truncated population.
    mappings = client.get(f"/api/jobs/{job_id}/mappings")
    if mappings.status_code == 200:
        body = mappings.json()
        decisions = body.get("decisions", body if isinstance(body, list) else [])
        assert not decisions
