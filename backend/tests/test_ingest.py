"""Ingestion: CSV/XLSX raw preservation, provenance, and parsing-issue detection."""
from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

from app.ingest import parse_file
from app.source_records import RawType, UnsupportedFormatError


def _parse(name, data):
    return parse_file(filename=name, data=data, stored_name="s" + name, max_bytes=10_000_000,
                      max_rows=1000)


def test_csv_preserves_leading_zeros_and_provenance():
    csv = b"EmployeeNumber,Full Name,Joined\n001,Alice,2021-03-15\n0027,Bob,2020-01-01\n"
    pf = _parse("legacy.csv", csv)
    assert len(pf.tables) == 1
    table = pf.tables[0]
    assert table.headers == ["EmployeeNumber", "Full Name", "Joined"]
    # leading zeros preserved as strings
    assert pf.records[0].cells[0].value == "001"
    assert pf.records[1].cells[0].value == "0027"
    # provenance: filename, table, row, col, header
    ref = pf.records[1].cells[0].ref
    assert ref.original_filename == "legacy.csv"
    assert ref.row_number == 2 and ref.col_index == 0 and ref.header == "EmployeeNumber"
    assert ref.sheet_name is None
    # future page/region provenance reserved but unused for CSV
    assert ref.page is None and ref.region is None


def test_csv_duplicate_header_detected_and_positional():
    csv = b"id,name,id\n1,a,2\n"
    pf = _parse("dup.csv", csv)
    kinds = {i.kind for i in pf.issues}
    assert "duplicate_header" in kinds
    # both id columns preserved by position (not merged)
    assert pf.tables[0].headers == ["id", "name", "id"]
    assert pf.records[0].cells[0].value == "1"
    assert pf.records[0].cells[2].value == "2"


def test_xlsx_multisheet_and_leading_zero_text():
    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"
    ws.append(["Emp ID", "Name", "Start"])
    ws.append(["001", "Alice", "2021-03-15"])
    ws.cell(row=2, column=1).number_format = "@"
    ws2 = wb.create_sheet("Contractors")
    ws2.append(["Emp ID", "Name"])
    ws2.append(["C01", "Bob"])
    buf = io.BytesIO()
    wb.save(buf)
    pf = _parse("export.xlsx", buf.getvalue())
    sheets = {t.sheet_name for t in pf.tables}
    assert sheets == {"Employees", "Contractors"}
    emp_records = [r for r in pf.records if r.ref.sheet_name == "Employees"]
    assert emp_records[0].cells[0].value == "001"  # text-preserved leading zeros
    assert emp_records[0].ref.sheet_name == "Employees"


def test_xlsx_formula_cell_flagged_not_executed():
    wb = Workbook()
    ws = wb.active
    ws.append(["a", "b"])
    ws.append([2, "=1+1"])
    buf = io.BytesIO()
    wb.save(buf)
    pf = _parse("f.xlsx", buf.getvalue())
    formula_issues = [i for i in pf.issues if i.kind == "formula_cell"]
    assert formula_issues, "formula cell must be flagged"
    # the formula text is preserved as-is; the cached value is not trusted
    cell = pf.records[0].cells[1]
    assert cell.raw_type == RawType.FORMULA
    assert cell.value == "=1+1"


def test_xlsx_numeric_identifier_zero_loss_flagged():
    wb = Workbook()
    ws = wb.active
    ws.append(["Emp ID", "Name"])
    ws.append([7, "Alice"])  # numeric id cell -> potential leading-zero loss
    buf = io.BytesIO()
    wb.save(buf)
    pf = _parse("n.xlsx", buf.getvalue())
    assert any(i.kind == "possible_zero_loss" for i in pf.issues)
    # value preserved as-is, never fabricated back into "007"
    assert pf.records[0].cells[0].value == "7"


def test_csv_quoting_preserves_cell_boundaries():
    # Quoted comma, escaped quote, and embedded newline must keep original cell boundaries.
    csv = b'id,note,name\r\n1,"a,b","say ""hi"""\r\n2,"line1\nline2",Bob\r\n'
    pf = _parse("q.csv", csv)
    assert pf.tables[0].headers == ["id", "note", "name"]
    assert pf.records[0].cells[1].value == "a,b"          # quoted comma stays one cell
    assert pf.records[0].cells[2].value == 'say "hi"'      # escaped quote unescaped
    assert pf.records[1].cells[1].value == "line1\nline2"  # embedded newline preserved
    assert len(pf.records) == 2


def test_unsupported_format_rejected():
    with pytest.raises(UnsupportedFormatError) as ei:
        _parse("resume.pdf", b"%PDF-1.4 ...")
    assert "CSV and XLSX" in str(ei.value)
    with pytest.raises(UnsupportedFormatError):
        _parse("data.json", b"{}")
