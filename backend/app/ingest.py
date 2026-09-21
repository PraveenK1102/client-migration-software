"""File ingestion: CSV and XLSX only (Milestone 1).

Kept strictly separate from the migration workflow. Parsers read raw values with
NO type coercion (so string leading zeros survive), preserve provenance, and
record problems as :class:`ParsingIssue` rather than guessing.

Unsupported formats (PDF, scanned images, handwriting, mixed-layout docs) are
rejected with an actionable message. Adding those is a documented future
extension, not implemented here — see README "Future extensions".
"""
from __future__ import annotations

import csv
import io
import re
import uuid
from pathlib import Path

from openpyxl import load_workbook

from .source_records import (
    ParsedFile,
    ParsingIssue,
    RawCell,
    RawType,
    RowLimitExceededError,
    SourceRecord,
    SourceRef,
    SourceTable,
    UnsupportedFormatError,
    UploadTooLargeError,
)

SUPPORTED_EXTENSIONS = {".csv", ".xlsx"}

_DATE_ISO = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_DATE_SLASH = re.compile(r"^\d{1,4}[/.]\d{1,2}[/.]\d{1,4}$")
_NUMBER = re.compile(r"^-?\d+(\.\d+)?$")
_LEADING_ZERO = re.compile(r"^0\d+$")


def _classify_text(value: str) -> RawType:
    """Light, non-coercing classification of a raw string cell."""
    v = value.strip()
    if v == "":
        return RawType.EMPTY
    if _DATE_ISO.match(v) or _DATE_SLASH.match(v):
        return RawType.DATE
    if v.lower() in {"true", "false", "yes", "no"}:
        return RawType.BOOL
    if _NUMBER.match(v):
        return RawType.NUMBER
    return RawType.TEXT


def validate_upload(filename: str, size_bytes: int, max_bytes: int) -> str:
    """Validate extension + size *before* parsing. Returns the lowercased extension."""
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFormatError(
            f"Unsupported file type '{ext or filename}'. Milestone 1 supports only "
            f"CSV and XLSX ({', '.join(sorted(SUPPORTED_EXTENSIONS))}). "
            f"PDF/scanned-document extraction is a documented future extension, not yet implemented."
        )
    if size_bytes > max_bytes:
        raise UploadTooLargeError(
            f"File '{filename}' is {size_bytes} bytes, exceeding the limit of {max_bytes} bytes."
        )
    return ext


def _dedupe_headers(headers: list[str], ref_base: SourceRef, issues: list[ParsingIssue]) -> list[str]:
    """Keep header identity by column position; report duplicates instead of overwriting."""
    seen: dict[str, int] = {}
    for idx, h in enumerate(headers):
        key = h.strip().lower()
        if key in seen:
            issues.append(
                ParsingIssue(
                    ref=ref_base.model_copy(update={"col_index": idx, "header": h}),
                    kind="duplicate_header",
                    detail=(
                        f"Duplicate header '{h}' at column {idx}; also at column {seen[key]}. "
                        f"Columns are preserved by position and not merged."
                    ),
                    severity="warning",
                )
            )
        else:
            seen[key] = idx
    return headers


def _row_limit_error(where: str, filename: str, max_rows: int) -> RowLimitExceededError:
    return RowLimitExceededError(
        f"{where} in '{filename}' has more than the configured limit of {max_rows} data rows. "
        f"The file was REJECTED rather than migrating a truncated subset of the employee "
        f"population. No rows from this file were ingested. To proceed, split the file into "
        f"parts within the limit, or raise MAX_ROWS_PER_TABLE for this deployment, then re-upload."
    )


def _parse_csv(
    data: bytes, file_id: str, filename: str, stored_name: str, max_rows: int
) -> ParsedFile:
    """Stream the CSV row-by-row (never ``list(reader)`` of the whole file) so an oversized table is
    rejected at limit+1 without materialising the remainder, and no partial rows are ever produced."""
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    issues: list[ParsingIssue] = []

    table_id = f"tbl_{uuid.uuid4().hex[:12]}"
    ref_base = SourceRef(file_id=file_id, original_filename=filename, table_id=table_id, sheet_name=None)

    headers: list[str] | None = None
    records: list[SourceRecord] = []
    phys = 0        # physical data-row ordinal after the header (incl. blanks) — provenance row_number
    kept = 0        # non-empty data rows actually ingested — counts toward the hard limit

    for raw_row in reader:
        if headers is None:
            # Skip leading fully-empty rows, then take the first non-empty row as headers.
            if all((c or "").strip() == "" for c in raw_row):
                continue
            headers = [(h or "").strip() for h in raw_row]
            _dedupe_headers(headers, ref_base, issues)
            continue
        # A fully-empty data row (e.g. a trailing blank line) is noise, not an employee: skip it
        # (never persisted, never counted) so it can neither be "lost" nor trip the limit.
        if all((c or "").strip() == "" for c in raw_row):
            phys += 1
            continue
        phys += 1
        kept += 1
        if kept > max_rows:
            raise _row_limit_error("CSV", filename, max_rows)
        row_ref = ref_base.model_copy(update={"row_number": phys})
        cells: list[RawCell] = []
        for c_i, header in enumerate(headers):
            raw_val = raw_row[c_i] if c_i < len(raw_row) else ""
            value = raw_val if raw_val != "" else None
            cells.append(
                RawCell(
                    col_index=c_i,
                    header=header,
                    value=value,
                    raw_type=_classify_text(raw_val),
                    ref=row_ref.model_copy(update={"col_index": c_i, "header": header}),
                )
            )
        records.append(SourceRecord(ref=row_ref, cells=cells))

    if headers is None:
        raise UnsupportedFormatError(f"CSV '{filename}' contains no readable rows.")

    table = SourceTable(
        table_id=table_id,
        file_id=file_id,
        original_filename=filename,
        sheet_name=None,
        headers=headers,
        n_rows=len(records),
    )
    return ParsedFile(
        file_id=file_id,
        original_filename=filename,
        content_type="text/csv",
        stored_name=stored_name,
        size_bytes=len(data),
        tables=[table],
        records=records,
        issues=issues,
    )


def _looks_like_identifier(header: str) -> bool:
    h = header.lower()
    return any(tok in h for tok in ("id", "no", "number", "code", "emp", "employee", "staff"))


def _parse_xlsx(
    data: bytes, file_id: str, filename: str, stored_name: str, max_rows: int
) -> ParsedFile:
    """Stream each worksheet in openpyxl read-only mode (a bounded workbook grid is never fully
    materialised) so an oversized sheet is rejected at limit+1 without loading the remainder.

    data_only=False so formula cells surface as formulas (we never trust cached values)."""
    wb = load_workbook(io.BytesIO(data), data_only=False, read_only=True)
    tables: list[SourceTable] = []
    records: list[SourceRecord] = []
    issues: list[ParsingIssue] = []

    try:
        for ws in wb.worksheets:
            table_id = f"tbl_{uuid.uuid4().hex[:12]}"
            ref_base = SourceRef(
                file_id=file_id, original_filename=filename, table_id=table_id, sheet_name=ws.title
            )
            headers: list[str] | None = None
            phys = 0        # physical data-row ordinal after the header (incl. blanks)
            kept = 0        # non-empty rows ingested — counts toward the hard limit

            for row_cells in ws.iter_rows(values_only=False):
                if headers is None:
                    # Skip leading empty rows, then take the first non-empty row as headers.
                    if all((c.value is None or str(c.value).strip() == "") for c in row_cells):
                        continue
                    headers = [
                        (str(c.value).strip() if c.value is not None else f"column_{i}")
                        for i, c in enumerate(row_cells)
                    ]
                    _dedupe_headers(headers, ref_base, issues)
                    continue
                if all(c.value is None or str(c.value).strip() == "" for c in row_cells):
                    phys += 1
                    continue  # skip blank rows entirely (never persisted, never counted)
                phys += 1
                kept += 1
                if kept > max_rows:
                    raise _row_limit_error(f"Sheet '{ws.title}'", filename, max_rows)
                row_ref = ref_base.model_copy(update={"row_number": phys})
                cells: list[RawCell] = []
                for c_i, header in enumerate(headers):
                    cell = row_cells[c_i] if c_i < len(row_cells) else None
                    cell_ref = row_ref.model_copy(update={"col_index": c_i, "header": header})
                    value, raw_type = _read_xlsx_cell(cell, header, cell_ref, issues)
                    cells.append(
                        RawCell(
                            col_index=c_i, header=header, value=value, raw_type=raw_type, ref=cell_ref
                        )
                    )
                records.append(SourceRecord(ref=row_ref, cells=cells))

            if headers is None:
                continue  # empty worksheet -> skip (only non-empty tables are ingested)
            tables.append(
                SourceTable(
                    table_id=table_id,
                    file_id=file_id,
                    original_filename=filename,
                    sheet_name=ws.title,
                    headers=headers,
                    n_rows=kept,
                )
            )
    finally:
        wb.close()

    if not tables:
        raise UnsupportedFormatError(f"XLSX '{filename}' contains no non-empty worksheets.")

    return ParsedFile(
        file_id=file_id,
        original_filename=filename,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        stored_name=stored_name,
        size_bytes=len(data),
        tables=tables,
        records=records,
        issues=issues,
    )


def _read_xlsx_cell(cell, header: str, cell_ref: SourceRef, issues: list[ParsingIssue]):
    """Read one openpyxl cell as a raw value + type, flagging formulas and possible zero loss."""
    if cell is None or cell.value is None:
        return None, RawType.EMPTY

    # Formula cells: never execute; flag as unsupported and keep the formula text as-is.
    if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
        issues.append(
            ParsingIssue(
                ref=cell_ref,
                kind="formula_cell",
                detail=(
                    f"Cell contains a formula ('{cell.value}'). Formulas are not evaluated in "
                    f"Milestone 1 and cached values are not treated as verified facts."
                ),
                severity="warning",
            )
        )
        return str(cell.value), RawType.FORMULA

    # Numbers: openpyxl already stripped any original leading zeros. If the column
    # looks like an identifier, the original zeros (if any) are unrecoverable — report it.
    if isinstance(cell.value, bool):
        return ("true" if cell.value else "false"), RawType.BOOL
    if isinstance(cell.value, (int, float)):
        text = repr(cell.value) if isinstance(cell.value, float) else str(cell.value)
        if isinstance(cell.value, int) and _looks_like_identifier(header):
            issues.append(
                ParsingIssue(
                    ref=cell_ref,
                    kind="possible_zero_loss",
                    detail=(
                        f"Identifier-like column '{header}' holds a numeric cell ({text}); if the "
                        f"original identifier had leading zeros they are unrecoverable from XLSX. "
                        f"Value is preserved as-is and NOT fabricated."
                    ),
                    severity="warning",
                )
            )
        return text, RawType.NUMBER

    # Dates/datetimes stored as real Excel dates.
    if hasattr(cell.value, "isoformat"):
        try:
            return cell.value.isoformat(), RawType.DATE
        except Exception:  # pragma: no cover - defensive
            return str(cell.value), RawType.DATE

    text = str(cell.value)
    return (text if text != "" else None), _classify_text(text)


def parse_file(
    *, filename: str, data: bytes, stored_name: str, max_bytes: int, max_rows: int,
    file_id: str | None = None,
) -> ParsedFile:
    """Parse an uploaded file into the typed source-record representation.

    Raises :class:`UnsupportedFormatError` / :class:`UploadTooLargeError` on bad input.
    """
    ext = validate_upload(filename, len(data), max_bytes)
    file_id = file_id or f"file_{uuid.uuid4().hex[:12]}"
    if ext == ".csv":
        return _parse_csv(data, file_id, filename, stored_name, max_rows)
    return _parse_xlsx(data, file_id, filename, stored_name, max_rows)
