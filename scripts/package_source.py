"""Package the COMPLETE current source into a review ZIP.

Allowlist-based (never a blind repo zip), secrets excluded and scanned, with a
SOURCE_MANIFEST.txt (relative paths + per-file SHA-256). Verifies the archive extracts
safely, contains no excluded/traversal paths, and holds no likely credentials before it
is considered attachable. If a secret is found, it ABORTS (does not write the zip).

Run:  python scripts/package_source.py
Output: artifacts/darwinbox_working_source_final.zip  (+ prints path, count, size, SHA-256)

This is the COMPLETE current working source for code inspection (it may include internal engineering
evidence and notes; it never includes secrets or runtime DBs). The clean, product-first reviewer
package is built separately by scripts/package_reviewer.py.
"""
from __future__ import annotations

import hashlib
import io
import re
import sys
import tempfile
import zipfile
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
ARTIFACTS = REPO / "artifacts"
ZIP_NAME = "darwinbox_working_source_final.zip"

# --- allowlist: explicit files + directory globs (relative to repo root) ---
ALLOW_FILES = [
    "README.md", "WRITEUP.md",
    ".gitignore",
    "backend/requirements.txt", "backend/requirements.lock.txt", "backend/pytest.ini",
    "backend/.env.example",
    "frontend/package.json", "frontend/package-lock.json", "frontend/tsconfig.json",
    "frontend/vite.config.ts", "frontend/index.html",
    "scripts/package_source.py", "scripts/package_reviewer.py", "scripts/run_target.py",
]
ALLOW_GLOBS = [
    "schemas/*.json", "schemas/*.yaml", "schemas/*.yml",
    "schemas/tenants/*.yaml", "schemas/tenants/*.yml", "schemas/tenants/*.json",
    "sample-data/*.py", "sample-data/*.csv", "sample-data/*.xlsx",
    "sample-data/final-demo/*.xlsx",       # the four cleaned final demo fixtures
    "sample-data/manual-demo/*.csv", "sample-data/manual-demo/*.xlsx", "sample-data/manual-demo/*.md",
    "sample-data/structured-demo/*.py", "sample-data/structured-demo/*.csv",
    "sample-data/structured-demo/*.xlsx", "sample-data/structured-demo/*.md",
    "reports/*.md", "reports/*.json",     # durable validation/stress/eval/security reports
    "backend/app/**/*.py",                # includes the restructured app/workflows/**
    "backend/mock_target/**/*.py",
    "backend/scripts/*.py",               # demo/observability/perf/generator scripts
    "backend/tests/**/*.py", "backend/tests/fixtures/*.json",
    "frontend/src/**/*",
]
# Hard exclusions (belt-and-suspenders even if an allowlist glob matches).
EXCLUDE_PARTS = {".git", ".venv", "node_modules", "dist", "__pycache__", ".pytest_cache",
                 "data", "artifacts", ".vite", "blind_packages", "external_packages"}
EXCLUDE_SUFFIX = (".db", ".db-wal", ".db-shm", ".sqlite", ".sqlite3", ".log", ".tsbuildinfo",
                  ".pyc", ".DS_Store")


def _excluded(rel: Path) -> bool:
    if any(part in EXCLUDE_PARTS for part in rel.parts):
        return True
    if rel.name == ".env" or (rel.name.endswith(".env") and rel.name != ".env.example"):
        return True
    if rel.suffix in EXCLUDE_SUFFIX or rel.name == ".DS_Store":
        return True
    return False


def collect() -> list[Path]:
    found: set[Path] = set()
    for f in ALLOW_FILES:
        p = REPO / f
        if p.is_file():
            found.add(p.resolve())
    for g in ALLOW_GLOBS:
        for p in REPO.glob(g):
            if p.is_file():
                found.add(p.resolve())
    result = []
    for p in sorted(found):
        rel = p.relative_to(REPO)
        if _excluded(rel):
            continue
        result.append(p)
    return result


# --- secret scanning (report path only, never matched content) ---
SECRET_PATTERNS = [
    re.compile(rb"gsk_[A-Za-z0-9]{20,}"),       # Groq live key
    re.compile(rb"sk-[A-Za-z0-9]{20,}"),        # generic provider key
    re.compile(rb"lsv2_[A-Za-z0-9_]{20,}"),     # LangSmith key (v2)
    re.compile(rb"ls__[A-Za-z0-9]{20,}"),       # LangSmith key (legacy)
    re.compile(rb"AKIA[0-9A-Z]{16}"),           # AWS access key id
    re.compile(rb"-----BEGIN [A-Z ]*PRIVATE KEY-----"),
]


def _configured_secret_bytes() -> list[bytes]:
    env = REPO / "backend" / ".env"
    out = []
    if env.exists():
        for line in env.read_text(encoding="utf-8", errors="ignore").splitlines():
            s = line.strip()
            for prefix in ("GROQ_API_KEY=", "LANGSMITH_API_KEY=", "LANGCHAIN_API_KEY="):
                if s.startswith(prefix) and s.split("=", 1)[1].strip():
                    out.append(s.split("=", 1)[1].strip().encode())
    return out


def scan_bytes(data: bytes, extra: list[bytes]) -> bool:
    for pat in SECRET_PATTERNS:
        if pat.search(data):
            return True
    for b in extra:
        if b and b in data:
            return True
    return False


def scan_xlsx_cells(path: Path, extra: list[bytes]) -> bool:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for cell in row:
                    if cell is None:
                        continue
                    if scan_bytes(str(cell).encode(), extra):
                        return True
        wb.close()
    except Exception:
        return False
    return False


def main() -> int:
    files = collect()
    extra = _configured_secret_bytes()
    flagged: list[str] = []
    for p in files:
        rel = str(p.relative_to(REPO))
        data = p.read_bytes()
        if p.suffix == ".xlsx":
            if scan_xlsx_cells(p, extra):
                flagged.append(rel)
            continue  # do not byte-scan the zip container of an xlsx (false positives)
        if scan_bytes(data, extra):
            flagged.append(rel)
    if flagged:
        print("SECRET SCAN FAILED — archive NOT written. Flagged (paths only):")
        for f in flagged:
            print("  -", f)
        return 2

    # Manifest (paths + per-file sha256), excluding the manifest itself.
    manifest_lines = ["# SOURCE_MANIFEST — darwinbox working source, final (relative path  sha256)"]
    for p in files:
        rel = str(p.relative_to(REPO))
        digest = hashlib.sha256(p.read_bytes()).hexdigest()
        manifest_lines.append(f"{rel}  {digest}")
    manifest = "\n".join(manifest_lines) + "\n"

    ARTIFACTS.mkdir(parents=True, exist_ok=True)
    zip_path = ARTIFACTS / ZIP_NAME
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for p in files:
            zf.write(p, str(p.relative_to(REPO)))
        zf.writestr("SOURCE_MANIFEST.txt", manifest)
    payload = buf.getvalue()

    # Verify: safe extraction (no traversal / excluded members), required files present.
    with zipfile.ZipFile(io.BytesIO(payload)) as zf:
        names = zf.namelist()
        for n in names:
            if n.startswith("/") or ".." in Path(n).parts:
                print(f"UNSAFE MEMBER PATH: {n}"); return 3
            if any(part in EXCLUDE_PARTS for part in Path(n).parts):
                print(f"EXCLUDED CONTENT LEAKED: {n}"); return 3
        with tempfile.TemporaryDirectory() as td:
            zf.extractall(td)  # ZipFile.extractall is path-safe in 3.12
    required = {"SOURCE_MANIFEST.txt", "backend/app/main.py", "backend/app/prepare.py",
                "backend/app/worker.py", "backend/app/blobstore.py", "backend/app/target_gateway.py",
                "backend/app/reconcile_target.py", "backend/app/versions.py",
                "backend/app/schema_loader.py", "backend/app/custom_fields.py", "backend/app/mapping_rules.py",
                "backend/app/delivery.py", "backend/app/policy.py",
                "backend/mock_target/service.py",
                # organization isolation + no-stuck delivery + eval + demo fixtures tests
                "backend/tests/test_m3g_organization_isolation.py",
                "backend/tests/test_m3b3_no_stuck_delivery.py",
                "backend/tests/test_eval_harness.py", "backend/tests/test_demo_fixtures.py",
                "backend/tests/conftest.py",
                # source intelligence + observability + eval engine
                "backend/app/date_inference.py", "backend/app/enum_inference.py",
                "backend/app/relationships.py", "backend/app/reference_integrity.py",
                "backend/app/transform_plan.py", "backend/app/source_intelligence.py",
                "backend/app/observability.py", "backend/app/eval_harness.py",
                "backend/app/model_projection.py", "backend/app/ingest.py", "backend/app/alias_audit.py",
                # restructured LangGraph workflows
                "backend/app/workflows/__init__.py", "backend/app/workflows/deps.py",
                "backend/app/workflows/_support.py",
                "backend/app/workflows/mapping/graph.py", "backend/app/workflows/mapping/state.py",
                "backend/app/workflows/mapping/context.py", "backend/app/workflows/mapping/edges.py",
                "backend/app/workflows/mapping/nodes/map_columns.py",
                "backend/app/workflows/preparation/graph.py",
                "backend/app/workflows/preparation/nodes/prepare.py",
                # the ONE final demo runner + the four cleaned fixtures
                "backend/scripts/final_four_migrations.py",
                "sample-data/final-demo/01_low_ambiguity_50x10.xlsx",
                "sample-data/final-demo/02_medium_ambiguity_20x10.xlsx",
                "sample-data/final-demo/03_high_ambiguity_10x10.xlsx",
                "sample-data/final-demo/04_clean_no_review_100x20.xlsx",
                # sample fixtures the test suite reads
                "sample-data/manual-demo/01_employee_master.csv",
                "sample-data/structured-demo/01_employees.csv", "sample-data/structured-demo/02_hr_details.xlsx",
                "sample-data/llm-generalization-demo.csv",
                "schemas/employee.v1.json", "schemas/employee.v2.yaml", "schemas/tenants/beta.yaml",
                # frontend
                "frontend/src/App.tsx", "frontend/src/components/DeliveryView.tsx",
                "frontend/src/components/MetricsView.tsx", "frontend/src/components/ReviewsView.tsx",
                "frontend/src/components/MigrationsView.tsx", "frontend/src/labels.ts",
                # reports + docs
                "reports/evaluation.md", "reports/evaluation.json",
                "scripts/package_reviewer.py",
                "README.md", "WRITEUP.md"}
    missing = [r for r in required if r not in names]
    if missing:
        print(f"MISSING REQUIRED MEMBERS: {missing}"); return 4

    zip_path.write_bytes(payload)
    sha = hashlib.sha256(payload).hexdigest()
    print("OK — source review bundle written")
    print(f"path: {zip_path}")
    print(f"files: {len(names)} (incl. SOURCE_MANIFEST.txt)")
    print(f"size_bytes: {len(payload)}")
    print(f"sha256: {sha}")
    print(f"manifest: SOURCE_MANIFEST.txt (inside the archive)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
